from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).parent
PPTX = OUT / "CrewPilot_AI_Track1_Pitch.pptx"
W, H = 13.333, 7.5

NAVY = "0B1622"
PANEL = "122332"
PANEL_2 = "182D3D"
INK = "102233"
WHITE = "F5F8FA"
PALE = "EAF0F5"
MUTED = "9EB1C1"
MUTED_DARK = "60798D"
CYAN = "55C8E0"
CYAN_DARK = "087E9C"
AMBER = "E8A23B"
AMBER_DARK = "86530B"
RED = "E97886"
GREEN = "65C89D"
LINE = "2B4354"
FONT_HEAD = "Aptos Display"
FONT_BODY = "Aptos"


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, line=None, radius=True, line_width=0.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(line_width)
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
         font=FONT_BODY, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
         margin=0, italic=False, tracking=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    if tracking is not None:
        r._r.get_or_add_rPr().set("spc", str(tracking))
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def circle(slide, x, y, d, fill, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color or fill)
    return shape


def base_slide(prs, dark=True, section="CREWPILOT  /  TRACK 1"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY if dark else PALE)
    text(slide, section, 0.62, 0.35, 7.0, 0.25, 9,
         CYAN if dark else CYAN_DARK, True, tracking=90)
    line(slide, 0.62, 7.08, 12.70, 7.08, LINE if dark else "D3DDE5", 0.7)
    text(slide, "THE INTELLIGENT FBO  ·  STAFFING", 0.62, 7.16, 6.0, 0.18,
         8, MUTED if dark else MUTED_DARK, True, tracking=65)
    text(slide, f"{len(prs.slides):02d}", 12.20, 7.13, 0.45, 0.22,
         9, MUTED if dark else MUTED_DARK, True, align=PP_ALIGN.RIGHT)
    return slide


def label(slide, value, x, y, w, dark=True, color=None):
    text(slide, value.upper(), x, y, w, 0.22, 9,
         color or (MUTED if dark else MUTED_DARK), True, tracking=90)


def section_title(slide, value, dark=True, y=0.82, size=30):
    text(slide, value, 0.62, y, 11.9, 0.62, size,
         WHITE if dark else INK, True, font=FONT_HEAD)


def add_notes():
    notes = """# CrewPilot AI — presenter notes

Six slides. About three minutes. The Track 1 brief does not set a pitch length; this is a compact rehearsal script.

## 1. Staff the rush
CrewPilot addresses Track 1, The Intelligent FBO. Our focus is staffing. The prototype helps an FBO manager compare a normal week with event demand, inspect staffing needs, and test disruptions. It is planning support. A person keeps final authority.

## 2. One arrival, several jobs
An aircraft may create ramp, fuel, and customer-service work. Fuel is only added when requested. The task times and crew counts shown are configurable demo assumptions. They are not validated FBO standards. When arrivals overlap, qualified coverage matters more than a single total headcount.

## 3. Plan the full week
The browser planner compares a synthetic ordinary week with three event cases across seven days. Low, medium, and high use fixed multipliers of 1.15, 1.4, and 1.7. Those are scenario inputs, not forecasts. The model translates estimated work into role and shift coverage.

## 4. Stress-test the plan
The simulator can add four unexpected aircraft, remove two ramp workers, and delay three expected arrivals by up to 45 minutes. It shows uncovered tasks under the disruption. The original plan and the recovery proposal face the same affected work.

## 5. Manager decides
The weekly planner builds shift coverage using peak hourly need, then assigns eligible fictional workers with a greedy rule. It checks the configured hours, rest, breaks, and qualifications. It is not a proven global optimizer. The manager reviews, approves, or rejects. Nothing is dispatched.

## 6. What comes next
The current demo uses synthetic arrivals, tasks, and workers. It has no validated Van Nuys or Los Angeles FBO demand data. So it cannot claim real-world forecast accuracy or proven operational savings. Next: validate the task model with an FBO operator, obtain approved historical demand, back-test forecasts, then measure uncovered work and labor tradeoffs.
"""
    (OUT / "Presenter_Notes.md").write_text(notes, encoding="utf-8")


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "CrewPilot AI — The Intelligent FBO"
    prs.core_properties.subject = "Track 1 staffing planning prototype"
    prs.core_properties.author = "CrewPilot AI"
    prs.core_properties.keywords = "NCBAA, SpartHill Hacks, staffing, FBO"

    # 1 — Cover: a single repeated shift-rail motif.
    s = base_slide(prs, True, "SPARTHILL HACKS  /  NCBAA  /  TRACK 1")
    text(s, "CREWPILOT", 0.72, 1.15, 4.7, 0.38, 13, CYAN, True, tracking=180)
    text(s, "Staff the rush.", 0.72, 1.78, 7.1, 0.88, 46, WHITE, True, FONT_HEAD)
    text(s, "Keep decisions human.", 0.72, 2.72, 7.1, 0.72, 35, AMBER, True, FONT_HEAD)
    text(s, "Seven-day FBO staffing under event demand.",
         0.75, 3.74, 6.6, 0.42, 19, "D0DCE5")
    text(s, "STAFFING PLANNING PROTOTYPE  ·  SYNTHETIC DEMO",
         0.75, 4.35, 6.6, 0.25, 10, MUTED, True, tracking=75)
    # Runway-like plan motif on the right, built from editable shapes.
    box(s, 8.48, 1.24, 4.0, 4.85, PANEL, LINE)
    label(s, "24-hour coverage", 8.82, 1.58, 3.3)
    text(s, "Three fixed shifts", 8.82, 1.95, 3.2, 0.4, 19, WHITE, True, FONT_HEAD)
    rail_x, rail_w = 8.98, 3.05
    for i, hour in enumerate(("00", "08", "16", "24")):
        x = rail_x + rail_w * i / 3
        line(s, x, 2.72, x, 5.43, LINE, 1)
        text(s, hour, x - 0.22, 2.44, 0.5, 0.2, 9, MUTED, True, align=PP_ALIGN.CENTER)
    shift_data = [("RAMP", 3.1, CYAN), ("FUEL", 3.75, AMBER), ("CUSTOMER", 4.4, GREEN)]
    for name, y, color in shift_data:
        text(s, name, 8.83, y + 0.03, 0.75, 0.22, 8, MUTED, True, tracking=55)
        for j in range(3):
            x = rail_x + j * rail_w / 3 + 0.04
            box(s, x, y, rail_w / 3 - 0.08, 0.34, color, color)
    text(s, "00–08", 9.12, 5.68, 0.72, 0.2, 9, MUTED, align=PP_ALIGN.CENTER)
    text(s, "08–16", 10.11, 5.68, 0.72, 0.2, 9, MUTED, align=PP_ALIGN.CENTER)
    text(s, "16–24", 11.10, 5.68, 0.72, 0.2, 9, MUTED, align=PP_ALIGN.CENTER)

    # 2 — Workload translation.
    s = base_slide(prs, False, "01  /  THE OPERATIONAL PROBLEM")
    section_title(s, "One arrival can trigger several jobs.", False, 0.88, 30)
    text(s, "Overlapping work needs qualified coverage by role.",
         0.65, 1.55, 8.8, 0.32, 16, MUTED_DARK)
    box(s, 0.72, 2.42, 2.55, 1.26, INK, INK)
    label(s, "Input", 0.98, 2.65, 1.8, True, CYAN)
    text(s, "Aircraft arrival", 0.98, 2.96, 2.0, 0.39, 19, WHITE, True, FONT_HEAD)
    text(s, "ETA + service request", 0.98, 3.38, 2.0, 0.2, 10, "C3D1DB")
    line(s, 3.42, 3.05, 4.10, 3.05, CYAN_DARK, 2.2)
    circle(s, 3.98, 2.94, 0.22, CYAN_DARK)
    cards = [
        ("RAMP", "2 crew", "35 min", CYAN_DARK),
        ("FUEL · IF REQUESTED", "1 crew", "40 min", AMBER_DARK),
        ("CUSTOMER SERVICE", "1 crew", "15 min", "25734E"),
    ]
    for i, (name, crew, duration, accent) in enumerate(cards):
        x = 4.42 + i * 2.75
        box(s, x, 2.40, 2.48, 1.32, WHITE, "D2DDE5")
        box(s, x, 2.40, 0.08, 1.32, accent, accent, radius=False)
        label(s, name, x + 0.25, 2.63, 2.0, False, accent)
        text(s, crew, x + 0.25, 2.95, 1.0, 0.32, 17, INK, True, FONT_HEAD)
        text(s, duration, x + 1.32, 2.98, 0.9, 0.26, 14, MUTED_DARK, True, align=PP_ALIGN.RIGHT)
    text(s, "EARLIEST START + DURATION · MIN AFTER ETA", 4.5, 4.12, 4.2, 0.22, 9, MUTED_DARK, True, tracking=40)
    task_x = 6.0
    task_rows = [("Customer", 0, 15, "25734E"), ("Ramp", 5, 35, CYAN_DARK), ("Fuel", 15, 40, AMBER_DARK)]
    for i, (name, offset, duration, color) in enumerate(task_rows):
        y = 4.48 + i * 0.39
        text(s, name, 4.5, y + 0.01, 1.28, 0.19, 9, MUTED_DARK, True)
        box(s, task_x + offset * 0.06, y, duration * 0.06, 0.2, color, color, radius=False)
    line(s, task_x, 5.70, task_x + 3.6, 5.70, "BCCAD5", 0.8)
    for minute in (0, 15, 30, 45, 60):
        x = task_x + minute * 0.06
        line(s, x, 5.66, x, 5.74, "9AABB8", 0.7)
        text(s, str(minute), x - 0.16, 5.77, 0.32, 0.16, 7, MUTED_DARK, align=PP_ALIGN.CENTER)
    box(s, 0.72, 6.14, 11.8, 0.48, "F8EEDC", "E8D2A7")
    text(s, "Demo task times and crew counts. Not validated FBO standards.",
         0.95, 6.25, 11.2, 0.2, 11, AMBER_DARK, True)

    # 3 — Seven-day comparison and demand cases.
    s = base_slide(prs, True, "02  /  SEVEN-DAY PREPARATION")
    section_title(s, "Compare a normal week with event demand.", True, 0.88, 29)
    text(s, "Same seeded profile. Three planning cases.", 0.65, 1.54, 7.3, 0.3, 15, MUTED)
    # Seven day tiles
    labels = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
    for i, day in enumerate(labels):
        x = 0.76 + i * 1.0
        box(s, x, 2.28, 0.82, 0.75, PANEL, LINE)
        text(s, day, x + 0.03, 2.51, 0.76, 0.2, 9, "DCE6ED", True, align=PP_ALIGN.CENTER, tracking=30)
        for dot in range(3):
            circle(s, x + 0.19 + dot * 0.17, 2.79, 0.07,
                   CYAN if i in (2, 3, 4) else "466071")
    label(s, "Normal baseline", 0.82, 3.45, 1.35)
    box(s, 2.3, 3.43, 2.2, 0.34, "416377", "416377", radius=False)
    text(s, "1.00×", 4.65, 3.45, 0.75, 0.22, 12, WHITE, True)
    label(s, "Event cases", 0.82, 4.12, 1.25)
    cases = [("LOW", "1.15×", 2.53, CYAN), ("MEDIUM", "1.40×", 3.08, AMBER), ("HIGH", "1.70×", 3.74, RED)]
    for i, (name, value, width, color) in enumerate(cases):
        y = 4.08 + i * 0.53
        text(s, name, 0.87, y + 0.05, 0.75, 0.2, 9, color, True, tracking=60)
        box(s, 2.3, y, width, 0.3, color, color, radius=False)
        text(s, value, 6.16, y + 0.04, 0.7, 0.2, 11, WHITE, True)
    box(s, 7.3, 2.25, 5.1, 3.25, PANEL, LINE)
    label(s, "The planner compares", 7.68, 2.62, 3.7)
    for i, (big, small, color) in enumerate([
        ("7 days", "ordinary vs. event", CYAN),
        ("3 cases", "low / medium / high", AMBER),
        ("Shift needs", "role by role", GREEN),
    ]):
        y = 3.17 + i * 0.69
        circle(s, 7.72, y + 0.05, 0.17, color)
        text(s, big, 8.08, y, 1.55, 0.3, 17, WHITE, True, FONT_HEAD)
        text(s, small, 9.9, y + 0.04, 2.0, 0.24, 12, MUTED)
    text(s, "Scenario multipliers are assumptions, not forecasts.",
         7.68, 5.02, 4.3, 0.22, 10, MUTED, italic=True)
    text(s, "1.15×  /  1.40×  /  1.70×", 0.82, 6.27, 4.2, 0.25, 10, MUTED, True, tracking=50)
    text(s, "SYNTHETIC DEMAND", 9.5, 6.27, 2.7, 0.25, 9, AMBER, True, align=PP_ALIGN.RIGHT, tracking=90)

    # 4 — Disruption simulator.
    s = base_slide(prs, True, "03  /  DISRUPTION TEST")
    section_title(s, "Stress-test the plan before the rush.", True, 0.88, 30)
    text(s, "A fixed demo scenario changes arrivals and crew availability.",
         0.65, 1.53, 9.8, 0.32, 15, MUTED)
    box(s, 0.75, 2.24, 3.0, 2.96, PANEL, LINE)
    label(s, "Start", 1.07, 2.55, 1.5)
    text(s, "Original\nstaffing plan", 1.07, 2.98, 2.3, 0.88, 25, WHITE, True, FONT_HEAD)
    text(s, "Hold assignments fixed", 1.07, 4.45, 2.3, 0.25, 11, MUTED)
    line(s, 3.98, 3.74, 4.70, 3.74, CYAN, 2.4)
    circle(s, 4.55, 3.60, 0.28, CYAN)
    box(s, 4.96, 2.24, 3.25, 2.96, PANEL_2, LINE)
    label(s, "Test", 5.28, 2.55, 1.5, color=AMBER)
    stressors = [("+4", "aircraft"), ("−2", "ramp crew"), ("+45m", "delays")]
    for i, (value, name) in enumerate(stressors):
        y = 3.03 + i * 0.58
        text(s, value, 5.33, y, 1.02, 0.34, 19, AMBER, True, FONT_HEAD)
        text(s, name, 6.42, y + 0.05, 1.25, 0.25, 13, WHITE)
    text(s, "Next 3 expected arrivals delayed", 5.3, 4.73, 2.6, 0.22, 10, MUTED)
    line(s, 8.44, 3.74, 9.12, 3.74, CYAN, 2.4)
    circle(s, 8.99, 3.60, 0.28, CYAN)
    box(s, 9.39, 2.24, 3.0, 2.96, "271D2A", "613642")
    label(s, "See the gaps", 9.72, 2.55, 2.2, color=RED)
    text(s, "Uncovered\nservice work", 9.72, 2.98, 2.4, 0.88, 24, WHITE, True, FONT_HEAD)
    text(s, "Then review a recovery plan", 9.72, 4.45, 2.3, 0.25, 11, "F2B4BC")
    box(s, 0.75, 5.78, 11.64, 0.62, PANEL, LINE)
    text(s, "Same disrupted demand", 1.02, 5.96, 2.35, 0.23, 12, CYAN, True)
    text(s, "Original plan", 4.15, 5.96, 1.5, 0.23, 11, WHITE, True)
    text(s, "↔", 5.75, 5.92, 0.35, 0.25, 14, MUTED, True, align=PP_ALIGN.CENTER)
    text(s, "Recovery proposal", 6.32, 5.96, 2.0, 0.23, 11, WHITE, True)
    text(s, "No preset improvement claim", 9.2, 5.96, 2.85, 0.23, 10, MUTED, align=PP_ALIGN.RIGHT)

    # 5 — Schedule + human decision.
    s = base_slide(prs, False, "04  /  PLAN  ·  REVIEW  ·  DECIDE")
    section_title(s, "Recommend shifts. Keep the decision human.", False, 0.88, 29)
    text(s, "Coverage, gaps, shifts, and labor cost—under the same demand.",
         0.65, 1.55, 10.2, 0.3, 15, MUTED_DARK)
    box(s, 0.75, 2.18, 7.1, 2.63, WHITE, "D1DCE5")
    label(s, "24-hour roster", 1.08, 2.48, 2.5, False)
    rails = [("RAMP", CYAN_DARK), ("FUEL", AMBER_DARK), ("CUSTOMER", "25734E")]
    for i, (role, color) in enumerate(rails):
        y = 3.08 + i * 0.48
        text(s, role, 1.1, y + 0.03, 0.85, 0.18, 8, MUTED_DARK, True, tracking=45)
        for shift in range(3):
            x = 2.05 + shift * 1.65
            box(s, x, y, 1.43, 0.28, color, color, radius=False)
    for i, label_text in enumerate(("00–08", "08–16", "16–24")):
        text(s, label_text, 2.16 + i * 1.65, 4.61, 1.2, 0.18, 9, MUTED_DARK, True, align=PP_ALIGN.CENTER)
    # Metrics are named but intentionally not assigned fabricated values.
    box(s, 8.18, 2.18, 4.2, 2.63, INK, INK)
    label(s, "Plan comparison", 8.53, 2.48, 2.8, True, CYAN)
    for i, metric in enumerate(("Work coverage", "Uncovered staff-hours", "Shift count + labor")):
        y = 3.02 + i * 0.49
        circle(s, 8.56, y + 0.055, 0.12, [CYAN, RED, AMBER][i])
        text(s, metric, 8.84, y, 2.95, 0.24, 13, WHITE)
    text(s, "Skills  ·  breaks  ·  hours  ·  rest", 8.54, 4.43, 3.4, 0.19, 9, MUTED, True)
    line(s, 6.95, 5.36, 6.95, 5.73, CYAN_DARK, 2)
    circle(s, 6.83, 5.6, 0.24, CYAN_DARK)
    box(s, 2.55, 5.93, 8.8, 0.66, INK, INK)
    text(s, "MANAGER REVIEW", 2.92, 6.15, 2.05, 0.22, 10, CYAN, True, tracking=75)
    text(s, "Approve  /  Reject", 5.44, 6.11, 2.65, 0.26, 16, WHITE, True, FONT_HEAD, align=PP_ALIGN.CENTER)
    text(s, "No dispatch. No aircraft control.", 8.6, 6.16, 2.3, 0.18, 9, MUTED, align=PP_ALIGN.RIGHT)

    # 6 — Honest evidence boundary and next proof.
    s = base_slide(prs, True, "05  /  WHAT WE KNOW  ·  WHAT WE NEED")
    section_title(s, "Useful prototype. Not a validated forecast.", True, 0.88, 30)
    box(s, 0.75, 2.03, 5.72, 3.56, PANEL, LINE)
    box(s, 6.78, 2.03, 5.72, 3.56, PANEL, LINE)
    label(s, "Works in the demo", 1.12, 2.4, 3.3, color=GREEN)
    current = ["Seven-day ordinary / event comparison", "Low, medium, high demand cases", "Shift gaps + disruption recovery", "Manager review and decision record"]
    for i, item in enumerate(current):
        y = 2.92 + i * 0.49
        circle(s, 1.15, y + 0.04, 0.13, GREEN)
        text(s, item, 1.48, y, 4.45, 0.25, 13, WHITE)
    label(s, "Still needs proof", 7.15, 2.4, 3.3, color=AMBER)
    missing = ["Observed LA FBO arrival + service data", "Forecast accuracy on real holdout data", "Operator review of tasks and labor rules", "Measured service and resource improvement"]
    for i, item in enumerate(missing):
        y = 2.92 + i * 0.49
        circle(s, 7.18, y + 0.04, 0.13, AMBER)
        text(s, item, 7.5, y, 4.55, 0.28, 13, WHITE)
    text(s, "SYNTHETIC DEMO · NO PASSENGER IDENTITY DATA", 1.12, 5.12, 4.9, 0.22,
         9, MUTED, True, tracking=65)
    text(s, "Next: validate demand → back-test → measure", 7.15, 5.12, 4.9, 0.22,
         10, MUTED, True)
    text(s, "AI assists. People decide.", 0.78, 6.0, 6.0, 0.55, 27, AMBER, True, FONT_HEAD)
    text(s, "CREWPILOT AI  /  STAFFING-ONLY PLANNING", 8.0, 6.17, 4.35, 0.22,
         9, MUTED, True, tracking=70, align=PP_ALIGN.RIGHT)

    prs.save(PPTX)
    add_notes()
    print(PPTX)


if __name__ == "__main__":
    make_deck()
